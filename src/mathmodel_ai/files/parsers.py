import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import polars as pl
import pymupdf
from charset_normalizer import __version__ as charset_normalizer_version
from charset_normalizer import from_bytes
from openpyxl import __version__ as openpyxl_version
from openpyxl import load_workbook
from PIL import Image
from PIL import __version__ as pillow_version
from pypdf import PdfReader
from pypdf import __version__ as pypdf_version

from mathmodel_ai.core.errors import FileParseError
from mathmodel_ai.schemas.files import ArtifactKind, FileKind, ImageMetadata, TextExtraction


@dataclass(frozen=True)
class ParsedDataset:
    name: str
    sheet_name: str | None
    frame: pl.DataFrame


@dataclass(frozen=True)
class GeneratedArtifact:
    kind: ArtifactKind
    name: str
    mime_type: str
    data: bytes
    metadata: dict[str, str | int | float | bool | None] = field(default_factory=dict)


@dataclass(frozen=True)
class ParserOutput:
    parser_name: str
    parser_version: str
    datasets: list[ParsedDataset] = field(default_factory=list)
    artifacts: list[GeneratedArtifact] = field(default_factory=list)
    text_extraction: TextExtraction | None = None
    image_metadata: ImageMetadata | None = None
    warnings: list[str] = field(default_factory=list)


class ParserRegistry:
    def __init__(self, *, max_pdf_pages: int, max_pdf_page_images: int) -> None:
        self._max_pdf_pages = max_pdf_pages
        self._max_pdf_page_images = max_pdf_page_images

    def parse(self, kind: FileKind, path: Path, *, name: str) -> ParserOutput:
        try:
            if kind is FileKind.CSV:
                return self._parse_csv(path, name=name)
            if kind is FileKind.EXCEL:
                return self._parse_xlsx(path, name=name)
            if kind is FileKind.PDF:
                return self._parse_pdf(path, name=name)
            if kind is FileKind.IMAGE:
                return self._parse_image(path)
            if kind is FileKind.TEXT:
                return self._parse_text(path)
        except FileParseError:
            raise
        except Exception as exc:
            raise FileParseError(f"{kind.value} parser failed safely") from exc
        raise FileParseError(f"no parser registered for {kind.value}")

    @staticmethod
    def _decoded_text(path: Path) -> str:
        match = from_bytes(path.read_bytes()).best()
        if match is None:
            raise FileParseError("text encoding could not be determined")
        return str(match)

    def _parse_csv(self, path: Path, *, name: str) -> ParserOutput:
        text = self._decoded_text(path)
        try:
            dialect = csv.Sniffer().sniff(text[:64_000], delimiters=",;\t|")
            frame = pl.read_csv(
                io.StringIO(text),
                separator=dialect.delimiter,
                try_parse_dates=True,
                infer_schema_length=10_000,
                null_values=["", "NA", "N/A", "null", "NULL"],
            )
        except Exception as exc:
            raise FileParseError("CSV content could not be parsed") from exc
        return ParserOutput(
            parser_name="polars_csv",
            parser_version=pl.__version__,
            datasets=[ParsedDataset(name=Path(name).stem, sheet_name=None, frame=frame)],
            artifacts=[self._preview_artifact(frame, f"{Path(name).stem}.preview.json")],
        )

    def _parse_xlsx(self, path: Path, *, name: str) -> ParserOutput:
        workbook = load_workbook(path, read_only=True, data_only=True)
        datasets: list[ParsedDataset] = []
        artifacts: list[GeneratedArtifact] = []
        warnings: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                header_row = next(rows, None)
                if header_row is None:
                    warnings.append(f"sheet {worksheet.title!r} is empty")
                    continue
                headers = self._headers(header_row)
                values = [dict(zip(headers, row, strict=False)) for row in rows]
                frame = pl.DataFrame(values, schema=headers, strict=False, infer_schema_length=None)
                dataset_name = f"{Path(name).stem}:{worksheet.title}"
                datasets.append(
                    ParsedDataset(
                        name=dataset_name,
                        sheet_name=worksheet.title,
                        frame=frame,
                    )
                )
                artifacts.append(
                    self._preview_artifact(
                        frame,
                        f"{self._safe_component(worksheet.title)}.preview.json",
                    )
                )
        finally:
            workbook.close()
        if not datasets:
            raise FileParseError("XLSX contains no non-empty worksheets")
        return ParserOutput(
            parser_name="openpyxl_polars",
            parser_version=f"openpyxl-{openpyxl_version}/polars-{pl.__version__}",
            datasets=datasets,
            artifacts=artifacts,
            warnings=warnings,
        )

    def _parse_pdf(self, path: Path, *, name: str) -> ParserOutput:
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise FileParseError("encrypted PDFs are not supported")
        if len(reader.pages) > self._max_pdf_pages:
            raise FileParseError("PDF page count exceeds configured limit")
        page_texts: list[str] = []
        pages_with_text = 0
        warnings: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages_with_text += 1
            else:
                warnings.append(f"page {page_number} has no extractable text")
            page_texts.append(text)
        full_text = "\n\n".join(page_texts)
        artifacts = [
            GeneratedArtifact(
                kind=ArtifactKind.EXTRACTED_TEXT,
                name=f"{Path(name).stem}.extracted.txt",
                mime_type="text/plain",
                data=full_text.encode("utf-8"),
                metadata={"page_count": len(reader.pages)},
            )
        ]
        document = pymupdf.open(path)  # type: ignore[no-untyped-call]
        try:
            for page_index in range(min(len(document), self._max_pdf_page_images)):
                page = document.load_page(page_index)  # type: ignore[no-untyped-call]
                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(1.25, 1.25),  # type: ignore[no-untyped-call]
                    alpha=False,
                )
                artifacts.append(
                    GeneratedArtifact(
                        kind=ArtifactKind.PDF_PAGE_IMAGE,
                        name=f"page-{page_index + 1:04d}.png",
                        mime_type="image/png",
                        data=pixmap.tobytes("png"),
                        metadata={"page_number": page_index + 1},
                    )
                )
        finally:
            document.close()  # type: ignore[no-untyped-call]
        return ParserOutput(
            parser_name="pypdf_pymupdf",
            parser_version=f"pypdf-{pypdf_version}/pymupdf-{pymupdf.__version__}",
            artifacts=artifacts,
            text_extraction=TextExtraction(
                text=full_text,
                page_count=len(reader.pages),
                pages_with_text=pages_with_text,
                warnings=warnings,
            ),
            warnings=warnings,
        )

    @staticmethod
    def _parse_image(path: Path) -> ParserOutput:
        with Image.open(path) as image:
            metadata = ImageMetadata(
                width=image.width,
                height=image.height,
                format=image.format or "unknown",
                mode=image.mode,
                frame_count=getattr(image, "n_frames", 1),
            )
        return ParserOutput(
            parser_name="pillow",
            parser_version=pillow_version,
            image_metadata=metadata,
        )

    def _parse_text(self, path: Path) -> ParserOutput:
        text = self._decoded_text(path)
        return ParserOutput(
            parser_name="charset_normalizer",
            parser_version=charset_normalizer_version,
            text_extraction=TextExtraction(text=text),
            artifacts=[
                GeneratedArtifact(
                    kind=ArtifactKind.EXTRACTED_TEXT,
                    name=f"{path.stem}.utf8.txt",
                    mime_type="text/plain",
                    data=text.encode("utf-8"),
                )
            ],
        )

    @staticmethod
    def _headers(values: tuple[Any, ...]) -> list[str]:
        headers: list[str] = []
        counts: dict[str, int] = {}
        for index, value in enumerate(values, start=1):
            base = (
                str(value).strip()
                if value is not None and str(value).strip()
                else f"column_{index}"
            )
            counts[base] = counts.get(base, 0) + 1
            headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
        return headers

    @staticmethod
    def _preview_artifact(frame: pl.DataFrame, name: str) -> GeneratedArtifact:
        payload = {
            "columns": frame.columns,
            "row_count": frame.height,
            "preview": frame.head(20).to_dicts(),
        }
        return GeneratedArtifact(
            kind=ArtifactKind.DATASET_PREVIEW,
            name=name,
            mime_type="application/json",
            data=json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8"),
            metadata={"preview_rows": min(frame.height, 20)},
        )

    @staticmethod
    def _safe_component(value: str) -> str:
        return "".join(
            character if character.isalnum() or character in "-_" else "_" for character in value
        )[:100]
